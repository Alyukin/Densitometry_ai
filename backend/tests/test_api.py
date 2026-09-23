import csv
import io
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.services.export import EXPORT_COLUMNS
from tests.conftest import upload, wait_done


def test_health(client: TestClient) -> None:
    r = client.get("/health")
    assert r.status_code == 200
    body = r.json()
    assert body["status"] == "ok"
    assert body["processor"] == "mock"
    assert body["is_mock"] is True
    assert body["database"] == "ok"


def test_openapi_has_required_endpoints(client: TestClient) -> None:
    paths = client.get("/openapi.json").json()["paths"]
    for p in [
        "/health",
        "/api/v1/studies/upload",
        "/api/v1/studies",
        "/api/v1/studies/{study_id}",
        "/api/v1/studies/{study_id}/process",
        "/api/v1/studies/{study_id}/status",
        "/api/v1/studies/{study_id}/result",
        "/api/v1/studies/{study_id}/download",
        "/api/v1/batch/process",
    ]:
        assert p in paths, p


def test_docs_pages(client: TestClient) -> None:
    r = client.get("/docs")
    assert r.status_code == 200
    assert "swagger-ui-bundle.js" in r.text
    r = client.get("/redoc")
    assert r.status_code == 200
    assert "redoc" in r.text.lower()


def test_upload_single_and_get(client: TestClient, samples: Path) -> None:
    r = upload(client, (samples / "study_spine_01/IM0001.dcm", "IM0001.dcm"))
    assert r.status_code == 201, r.text
    body = r.json()
    assert len(body["studies"]) == 1 and body["rejected"] == []
    study = body["studies"][0]
    assert study["status"] == "uploaded"
    assert study["image_count"] == 1
    assert study["modality"] == "BMD"

    detail = client.get(f"/api/v1/studies/{study['id']}").json()
    assert detail["images"][0]["body_part_examined"] == "LSPINE"
    assert detail["images"][0]["preview_url"]
    assert "PatientName" not in str(detail)

    listing = client.get("/api/v1/studies").json()
    assert listing["total"] == 1


def test_upload_groups_by_study_uid(client: TestClient, samples: Path) -> None:
    r = upload(
        client,
        (samples / "study_hip_01/IM0001.dcm", "study_hip_01/IM0001.dcm"),
        (samples / "study_hip_01/IM0002.dcm", "study_hip_01/IM0002.dcm"),
        (samples / "study_spine_01/IM0001.dcm", "study_spine_01/IM0001.dcm"),
    )
    assert r.status_code == 201
    studies = r.json()["studies"]
    assert sorted(s["image_count"] for s in studies) == [1, 2]
    hip = next(s for s in studies if s["image_count"] == 2)
    assert hip["name"] == "study_hip_01"
    assert hip["source_path"] == "study_hip_01"


def test_upload_zip(client: TestClient, samples: Path) -> None:
    r = upload(client, (samples / "demo_studies.zip", "demo_studies.zip"))
    assert r.status_code == 201, r.text
    studies = r.json()["studies"]
    assert len(studies) == 7
    assert sum(s["image_count"] for s in studies) == 10


def test_non_dicom_file_becomes_a_failure_row(client: TestClient, samples: Path) -> None:
    """Заказчик: файл не открывается или не разбирается как DICOM — processing_status = Failure.

    Значит, такой файл не отбрасывается при загрузке, а получает строку в выгрузке.
    """
    r = upload(client, (samples / "edge_cases/not_a_dicom.dcm", "not_a_dicom.dcm"))
    assert r.status_code == 201, r.text
    assert r.json()["rejected"] == []
    assert any("Failure" in w for w in r.json()["warnings"])
    sid = r.json()["studies"][0]["id"]
    client.post(f"/api/v1/studies/{sid}/process")
    wait_done(client, sid)
    row = client.get(f"/api/v1/studies/{sid}/result").json()["rows"][0]
    assert row["processing_status"] == "error"
    assert "не является DICOM" in row["error_message"]

    text = client.get(f"/api/v1/studies/{sid}/download?format=csv").content.decode("utf-8")
    line = text.splitlines()[1].split(",")
    assert line[0] == "not_a_dicom.dcm"  # path_to_study
    assert "Failure" in line


def test_broken_file_joins_the_study_from_its_folder(client: TestClient, samples: Path) -> None:
    r = upload(
        client,
        (samples / "study_spine_01/IM0001.dcm", "study/IM0001.dcm"),
        (samples / "edge_cases/not_a_dicom.dcm", "study/IM0002.dcm"),
    )
    studies = r.json()["studies"]
    assert len(studies) == 1 and studies[0]["image_count"] == 2
    sid = studies[0]["id"]
    client.post(f"/api/v1/studies/{sid}/process")
    wait_done(client, sid)
    rows = client.get(f"/api/v1/studies/{sid}/result").json()["rows"]
    assert sorted(r["processing_status"] for r in rows) == ["error", "success"]
    assert {r["path_to_study"] for r in rows} == {"study"}


def test_empty_file_is_a_failure_row_but_documents_and_junk_are_skipped(client: TestClient, samples: Path) -> None:
    empty = samples.parent / "empty.dcm"
    empty.write_bytes(b"")
    r = upload(
        client,
        (empty, "empty.dcm"),
        (samples / "edge_cases/not_a_dicom.dcm", "разметка.xlsx"),
        (samples / "edge_cases/not_a_dicom.dcm", ".DS_Store"),
    )
    assert r.status_code == 201
    assert [x["filename"] for x in r.json()["rejected"]] == ["разметка.xlsx"]
    assert len(r.json()["studies"]) == 1  # только пустой файл, служебный пропущен молча


def test_upload_duplicate_sop_rejected(client: TestClient, samples: Path) -> None:
    p = samples / "study_spine_01/IM0001.dcm"
    r = upload(client, (p, "a.dcm"), (p, "b.dcm"))
    assert r.status_code == 201
    assert r.json()["studies"][0]["image_count"] == 1
    assert "Дубликат" in r.json()["rejected"][0]["reason"]


def test_reupload_of_the_same_files_does_not_duplicate_the_study(client: TestClient, samples: Path) -> None:
    files = [(samples / f"study_hip_01/IM000{i}.dcm", f"study_hip_01/IM000{i}.dcm") for i in (1, 2)]
    sid = upload(client, *files).json()["studies"][0]["id"]
    client.post(f"/api/v1/studies/{sid}/process")
    wait_done(client, sid)

    r = upload(client, *files)
    assert r.status_code == 422
    assert "уже загружены" in r.json()["detail"]["message"]
    assert all(x["reason"].startswith("Уже загружен") for x in r.json()["detail"]["rejected"])
    assert client.get("/api/v1/studies").json()["total"] == 1
    csv_rows = client.get("/api/v1/batch/download?format=csv").text.strip().splitlines()[1:]
    assert len(csv_rows) == 2  # по строке на снимок, без задвоения


def test_study_uploaded_in_parts_stays_one_study(client: TestClient, samples: Path) -> None:
    first = upload(client, (samples / "study_hip_01/IM0001.dcm", "study_hip_01/IM0001.dcm")).json()["studies"][0]
    client.post(f"/api/v1/studies/{first['id']}/process")
    wait_done(client, first["id"])

    r = upload(client, (samples / "study_hip_01/IM0002.dcm", "study_hip_01/IM0002.dcm"))
    assert r.status_code == 201, r.text
    [study] = r.json()["studies"]
    assert study["id"] == first["id"]
    assert study["image_count"] == 2
    assert study["status"] == "uploaded"  # новый снимок ещё не обработан
    assert any("Добавлено снимков" in w for w in r.json()["warnings"])
    assert client.get("/api/v1/studies").json()["total"] == 1

    client.post(f"/api/v1/studies/{first['id']}/process")
    assert wait_done(client, first["id"])["status"] == "completed"
    assert len(client.get(f"/api/v1/studies/{first['id']}/result").json()["rows"]) == 2


def test_images_are_not_added_to_a_study_being_processed(client: TestClient, samples: Path) -> None:
    from app.db.session import session_scope
    from app.models import Study, StudyStatus

    sid = upload(client, (samples / "study_hip_01/IM0001.dcm", "IM0001.dcm")).json()["studies"][0]["id"]
    with session_scope() as db:
        db.get(Study, sid).status = StudyStatus.processing
    r = upload(client, (samples / "study_hip_01/IM0002.dcm", "IM0002.dcm"))
    assert r.status_code == 422
    assert "обрабатывается" in r.json()["detail"]["rejected"][0]["reason"]
    assert client.get(f"/api/v1/studies/{sid}").json()["image_count"] == 1


def test_full_processing_flow(client: TestClient, samples: Path) -> None:
    files = [(samples / f"study_combined/IM000{i}.dcm", f"study_combined/IM000{i}.dcm") for i in (1, 2, 3)]
    sid = upload(client, *files).json()["studies"][0]["id"]

    assert client.get(f"/api/v1/studies/{sid}/result").status_code == 409
    assert client.get(f"/api/v1/studies/{sid}/download").status_code == 409

    r = client.post(f"/api/v1/studies/{sid}/process")
    assert r.status_code == 202
    assert r.json()["status"] == "queued"

    st = wait_done(client, sid)
    assert st["status"] == "completed"
    assert st["progress"] == 100
    assert st["processed_images"] == st["total_images"] == 3

    res = client.get(f"/api/v1/studies/{sid}/result").json()
    assert res["is_mock"] is True
    assert len(res["rows"]) == 3
    regions = {row["anatomical_region"] for row in res["rows"]}
    assert regions == {"lumbar_spine", "proximal_femur"}
    for row in res["rows"]:
        assert row["processing_status"] == "success"
        assert row["quality_class"] in ("acceptable", "unacceptable")
        assert (row["quality_class"] == "unacceptable") == bool(row["violation_type"])
        assert row["details"]["checks"]

    csv_resp = client.get(f"/api/v1/studies/{sid}/download?format=csv")
    assert csv_resp.status_code == 200
    assert "attachment" in csv_resp.headers["content-disposition"]
    rows = list(csv.reader(io.StringIO(csv_resp.text)))
    assert rows[0] == EXPORT_COLUMNS
    assert len(rows) == 4  # шапка + 3 снимка исследования

    xlsx_resp = client.get(f"/api/v1/studies/{sid}/download?format=xlsx")
    assert xlsx_resp.status_code == 200
    wb = load_workbook(io.BytesIO(xlsx_resp.content))
    header = [c.value for c in wb["results"][1]]
    assert header == EXPORT_COLUMNS
    assert wb["results"].max_row == 4

    assert client.get(f"/api/v1/studies/{sid}/download?format=pdf").status_code == 422


def test_results_are_reproducible(client: TestClient, samples: Path) -> None:
    sid = upload(client, (samples / "demo_studies.zip", "demo.zip")).json()["studies"][0]["id"]
    client.post(f"/api/v1/studies/{sid}/process")
    wait_done(client, sid)
    first = client.get(f"/api/v1/studies/{sid}/result").json()["rows"]
    client.post(f"/api/v1/studies/{sid}/process")
    wait_done(client, sid)
    second = client.get(f"/api/v1/studies/{sid}/result").json()["rows"]
    key = ("image_uid", "anatomical_region", "quality_class", "violation_type")
    assert [[r[k] for k in key] for r in first] == [[r[k] for k in key] for r in second]


def test_dicom_without_pixels_is_non_standard_not_failure(client: TestClient, samples: Path) -> None:
    """DICOM открывается, но изображения в нём нет: это не Failure, а нестандартные данные."""
    sid = upload(client, (samples / "edge_cases/no_pixel_data.dcm", "nopix.dcm")).json()["studies"][0]["id"]
    client.post(f"/api/v1/studies/{sid}/process")
    st = wait_done(client, sid)
    assert st["status"] == "completed"
    row = client.get(f"/api/v1/studies/{sid}/result").json()["rows"][0]
    assert row["processing_status"] == "success"
    assert row["anatomical_region"] is None and row["quality_class"] is None
    assert "нет изображения" in row["details"]["non_standard"]
    assert row["error_message"] is None


def test_batch_process_and_download(client: TestClient, samples: Path) -> None:
    ids = [s["id"] for s in upload(client, (samples / "demo_studies.zip", "demo.zip")).json()["studies"]]

    assert client.post("/api/v1/batch/process", json={}).status_code == 400
    assert client.get("/api/v1/batch/download").status_code == 409

    r = client.post("/api/v1/batch/process", json={"study_ids": [ids[0], "missing"]})
    assert r.status_code == 202
    assert r.json()["accepted"] == [ids[0]]
    assert r.json()["skipped"][0] == {"study_id": "missing", "reason": "not_found"}
    wait_done(client, ids[0])

    r = client.post("/api/v1/batch/process", json={"all_pending": True})
    assert sorted(r.json()["accepted"]) == sorted(ids[1:])
    for sid in ids:
        assert wait_done(client, sid)["status"] == "completed"

    r = client.post("/api/v1/batch/process", json={"study_ids": ids})
    assert r.json()["accepted"] == []
    assert {s["reason"] for s in r.json()["skipped"]} == {"already_completed"}

    r = client.get("/api/v1/batch/download", params={"format": "csv"})
    assert r.status_code == 200
    assert len(r.text.strip().splitlines()) == 1 + 10

    r = client.get("/api/v1/batch/download", params={"format": "xlsx", "study_ids": ids[:1]})
    assert r.status_code == 200

    listing = client.get("/api/v1/studies", params={"status": "completed"}).json()
    assert listing["total"] == 7
    assert listing["items"][0]["summary"]["total"] >= 1


def test_preview_and_delete(client: TestClient, samples: Path) -> None:
    study = upload(client, (samples / "study_spine_01/IM0001.dcm", "IM0001.dcm")).json()["studies"][0]
    detail = client.get(f"/api/v1/studies/{study['id']}").json()
    r = client.get(detail["images"][0]["preview_url"])
    assert r.status_code == 200
    assert r.headers["content-type"] == "image/png"
    assert r.content[:8] == b"\x89PNG\r\n\x1a\n"

    assert client.delete(f"/api/v1/studies/{study['id']}").status_code == 204
    assert client.get(f"/api/v1/studies/{study['id']}").status_code == 404
    assert client.post(f"/api/v1/studies/{study['id']}/process").status_code == 404
