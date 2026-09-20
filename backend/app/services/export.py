"""CSV / XLSX export in the format required by the technical specification."""

from __future__ import annotations

import csv
import io
from collections.abc import Iterable
from datetime import UTC, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models import ImageResult

# Порядок колонок строго по ТЗ; quality_prob разрешён заказчиком (ответ на вопрос 8)
EXPORT_COLUMNS = [
    "path_to_study",
    "study_uid",
    "image_uid",
    "anatomical_region",
    "quality_class",
    "violation_type",
    "processing_status",
    "time_of_processing",
    "quality_prob",
]

CONTENT_TYPES = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


# По ТЗ (п. 2.5): quality_class — Integer 0/1, processing_status — Success / Failure.
# Внутри сервиса статусы хранятся в нижнем регистре и различают ошибку и таймаут;
# в итоговый файл они сводятся к двум значениям из ТЗ.
STATUS_TZ = {"success": "Success"}


def _quality_class(value: str | None) -> int | str:
    """0 / 1 целым числом; пустая строка, если область не оценена."""
    if value is None or value == "":
        return ""
    return 1 if str(value).strip().lower() in {"1", "unacceptable", "bad", "poor"} else 0


# Что попадает в колонки ТЗ: «auto» — то, что решил сервис, «reviewed» — то, что
# утвердил врач. По умолчанию auto, и это принципиально: по этому файлу оценивают
# модель, а в нём не должно оказаться исправлений, сделанных человеком.
SOURCE_AUTO = "auto"
SOURCE_REVIEWED = "reviewed"


def result_to_row(r: ImageResult, source: str = SOURCE_AUTO) -> dict:
    quality = r.final_quality_class if source == SOURCE_REVIEWED else r.quality_class
    violations = r.final_violation_type if source == SOURCE_REVIEWED else r.violation_type
    return {
        "path_to_study": r.path_to_study,
        "study_uid": r.study_uid or "",
        "image_uid": r.image_uid or "",
        "anatomical_region": r.anatomical_region or "",
        "quality_class": _quality_class(quality),
        "violation_type": violations or "",
        "processing_status": STATUS_TZ.get(r.processing_status, "Failure"),
        "time_of_processing": round(r.time_of_processing or 0.0, 3),
        "quality_prob": "" if r.confidence is None else round(float(r.confidence), 4),
    }


def to_csv(results: Iterable[ImageResult], source: str = SOURCE_AUTO) -> bytes:
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=EXPORT_COLUMNS, lineterminator="\n")
    writer.writeheader()
    for r in results:
        writer.writerow(result_to_row(r, source))
    return buf.getvalue().encode("utf-8")


def _autosize(ws) -> None:  # noqa: ANN001
    for col_cells in ws.columns:
        width = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(max(width + 2, 10), 80)


def to_xlsx(results: Iterable[ImageResult], is_mock: bool = False, source: str = SOURCE_AUTO) -> bytes:
    results = list(results)
    wb = Workbook()
    ws = wb.active
    ws.title = "results"
    ws.append(EXPORT_COLUMNS)
    header_fill = PatternFill("solid", fgColor="1F4E79")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for r in results:
        row = result_to_row(r, source)
        ws.append([row[c] for c in EXPORT_COLUMNS])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    _autosize(ws)

    # Лист объяснений: по строке на проверку — измеренная величина, порог, источник
    ws2 = wb.create_sheet("checks")
    ws2.append(
        [
            "image_uid",
            "check",
            "измерено",
            "критерий",
            "нарушение",
            "источник",
            "порог ТЗ",
            "по букве ТЗ",
            "участвует в вердикте",
            "error",
        ]
    )
    for c in ws2[1]:
        c.font = Font(bold=True)
    for r in results:
        checks = (r.details or {}).get("checks") or []
        if not checks:
            ws2.append([r.image_uid, "", "", "", "", "", "", "", "", r.error_message])
        for ch in checks:
            value = ch.get("value")
            tz_fired = ch.get("tz_fired")
            ws2.append(
                [
                    r.image_uid,
                    ch.get("rule_id") or ch.get("code"),
                    ch.get("measured") or ch.get("title"),
                    ch.get("criterion") or "",
                    "да" if (ch.get("fired") if "fired" in ch else ch.get("passed") is False) else "нет",
                    ch.get("source") or "",
                    ch.get("tz_threshold") if ch.get("tz_threshold") is not None else "",
                    "" if tz_fired is None else ("превышен" if tz_fired else "в норме"),
                    "да" if ch.get("decides", True) else "нет (справочно)",
                    r.error_message,
                ]
            )
            if isinstance(value, dict):
                ws2.cell(ws2.max_row, 3).comment = None
    _autosize(ws2)

    # Лист проверки специалистом: автоматический вердикт и решение врача рядом,
    # чтобы было видно, где сервис ошибся, и чтобы правки человека не смешивались
    # с тем, что выдала модель.
    ws3 = wb.create_sheet("review")
    ws3.append(
        [
            "image_uid",
            "область",
            "quality_class (сервис)",
            "violation_type (сервис)",
            "проверка",
            "quality_class (врач)",
            "violation_type (врач)",
            "проверил",
            "когда",
            "комментарий",
        ]
    )
    for c in ws3[1]:
        c.font = Font(bold=True)
    status_ru = {"": "не проверено", "confirmed": "подтверждено", "corrected": "исправлено"}
    for r in results:
        reviewed = r.review_status in ("confirmed", "corrected")
        ws3.append(
            [
                r.image_uid,
                r.anatomical_region or "",
                _quality_class(r.quality_class),
                r.violation_type or "",
                status_ru.get(r.review_status or "", r.review_status or ""),
                _quality_class(r.reviewed_quality_class) if reviewed else "",
                (r.reviewed_violation_type or "") if reviewed else "",
                r.reviewed_by or "",
                r.reviewed_at.isoformat(timespec="seconds") if r.reviewed_at else "",
                r.review_comment or "",
            ]
        )
    _autosize(ws3)

    meta = wb.create_sheet("meta")
    meta.append(["generated_at", datetime.now(UTC).isoformat(timespec="seconds")])
    meta.append(["rows", len(results)])
    meta.append(
        [
            "source",
            "решение специалиста, где оно есть" if source == SOURCE_REVIEWED else "автоматический вердикт сервиса",
        ]
    )
    meta.append(["mock", "yes — тестовые результаты, не для клинического использования" if is_mock else "no"])
    _autosize(meta)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
