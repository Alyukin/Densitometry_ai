"""Пункты ТЗ из раздела «Дополнительно»: DICOM SR, вторичная серия, проверка специалистом.

Плюс два свойства, которые ТЗ требует в технических требованиях и которые легко
потерять незаметно: воспроизводимость результата и независимость решения от
персональных данных.
"""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

import pydicom
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from tests.conftest import upload, wait_done

COMPREHENSIVE_SR = "1.2.840.10008.5.1.4.1.1.88.33"
SECONDARY_CAPTURE = "1.2.840.10008.5.1.4.1.1.7"


@pytest.fixture()
def processed(rb_client: TestClient, samples: Path) -> tuple[str, list[dict]]:
    """Обработанное исследование: снимок позвоночника и снимок бедра."""
    files = sorted((samples / "study_combined").glob("*.dcm"))
    sid = upload(rb_client, *[(f, f.name) for f in files]).json()["studies"][0]["id"]
    rb_client.post(f"/api/v1/studies/{sid}/process")
    assert wait_done(rb_client, sid)["status"] == "completed"
    rows = rb_client.get(f"/api/v1/studies/{sid}/result").json()["rows"]
    return sid, rows


def _items(sequence, value_type: str) -> list:
    return [i for i in sequence if i.ValueType == value_type]


def _texts(ds) -> dict[str, str]:
    """Все TEXT-элементы отчёта: смысл понятия -> значение."""
    out: dict[str, str] = {}

    def walk(seq):
        for item in seq:
            if item.ValueType == "TEXT":
                out[item.ConceptNameCodeSequence[0].CodeMeaning] = item.TextValue
            if "ContentSequence" in item:
                walk(item.ContentSequence)

    walk(ds.ContentSequence)
    return out


# --- DICOM SR ---------------------------------------------------------------


def test_sr_is_a_valid_comprehensive_report(rb_client: TestClient, processed) -> None:
    sid, _ = processed
    r = rb_client.get(f"/api/v1/studies/{sid}/sr")
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/dicom"

    ds = pydicom.dcmread(io.BytesIO(r.content))
    assert ds.SOPClassUID == COMPREHENSIVE_SR
    assert ds.Modality == "SR"
    assert ds.ValueType == "CONTAINER"
    assert ds.CompletionFlag == "COMPLETE"
    assert ds.SpecificCharacterSet == "ISO_IR 192"  # заключение на русском


def test_sr_carries_the_whole_verdict(rb_client: TestClient, processed) -> None:
    sid, rows = processed
    ds = pydicom.dcmread(io.BytesIO(rb_client.get(f"/api/v1/studies/{sid}/sr").content))
    texts = _texts(ds)
    assert texts["Анатомическая область"] in ("Поясничный отдел позвоночника", "Проксимальный отдел бедра")
    assert texts["Класс качества (0 — годно, 1 — нарушение)"] in ("0", "1")
    assert "Тип нарушения" in texts
    assert "не медицинское заключение" in texts["Ограничение"]

    containers = _items(ds.ContentSequence, "CONTAINER")
    assert len(containers) == len(rows)
    # у каждого снимка — ссылка на его SOP Instance и измеренные величины
    for c in containers:
        assert _items(c.ContentSequence, "IMAGE")
        nums = _items(c.ContentSequence, "NUM")
        assert nums, "измеренные величины должны попадать в отчёт"
        m = nums[0].MeasuredValueSequence[0]
        assert m.MeasurementUnitsCodeSequence[0].CodingSchemeDesignator == "UCUM"


def test_sr_keeps_found_structures_as_coordinates(rb_client: TestClient, processed) -> None:
    sid, _ = processed
    ds = pydicom.dcmread(io.BytesIO(rb_client.get(f"/api/v1/studies/{sid}/sr").content))
    scoords = [s for c in _items(ds.ContentSequence, "CONTAINER") for s in _items(c.ContentSequence, "SCOORD")]
    assert scoords, "оси и ориентиры должны выгружаться координатами"
    for s in scoords:
        assert s.GraphicType in ("POINT", "POLYLINE")
        assert len(s.GraphicData) % 2 == 0


def test_sr_ties_itself_to_the_source_study(rb_client: TestClient, processed, samples: Path) -> None:
    """Без идентификаторов исследования отчёт не ляжет в PACS к нужному пациенту."""
    sid, _ = processed
    ds = pydicom.dcmread(io.BytesIO(rb_client.get(f"/api/v1/studies/{sid}/sr").content))
    source = pydicom.dcmread(next((samples / "study_combined").glob("*.dcm")), stop_before_pixels=True)
    assert ds.StudyInstanceUID == source.StudyInstanceUID
    assert ds.SeriesInstanceUID != source.SeriesInstanceUID  # своя серия, а не поверх исходной


def test_sr_requires_results(rb_client: TestClient, samples: Path) -> None:
    f = next((samples / "study_spine_01").glob("*.dcm"))
    sid = upload(rb_client, (f, f.name)).json()["studies"][0]["id"]
    assert rb_client.get(f"/api/v1/studies/{sid}/sr").status_code == 409


# --- ZIP-пакет с дополнительными сериями --------------------------------------


def test_package_has_sr_overlay_series_and_tables(rb_client: TestClient, processed) -> None:
    sid, rows = processed
    r = rb_client.get(f"/api/v1/studies/{sid}/package")
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/zip"

    with zipfile.ZipFile(io.BytesIO(r.content)) as z:
        names = z.namelist()
        assert {"README.txt", "results.csv", "results.xlsx", "sr.dcm"} <= set(names)
        sc_names = [n for n in names if n.startswith("overlay/") and n.endswith(".dcm")]
        assert len(sc_names) == len(rows)
        assert len(set(sc_names)) == len(sc_names), "имена в архиве не должны повторяться"

        sc = pydicom.dcmread(io.BytesIO(z.read(sc_names[0])))
        assert sc.SOPClassUID == SECONDARY_CAPTURE
        assert sc.Modality == "OT"
        assert sc.PhotometricInterpretation == "RGB"
        assert sc.Rows > 0 and sc.Columns > 0
        assert len(sc.PixelData) == sc.Rows * sc.Columns * 3


# --- проверка специалистом ----------------------------------------------------


def test_confirm_keeps_the_automatic_verdict_intact(rb_client: TestClient, processed) -> None:
    sid, rows = processed
    row = rows[0]
    r = rb_client.post(
        f"/api/v1/studies/{sid}/images/{row['image_id']}/review",
        json={"action": "confirm", "reviewed_by": "Иванов И.И."},
    )
    assert r.status_code == 200
    out = r.json()["row"]
    assert out["review_status"] == "confirmed"
    assert out["reviewed_by"] == "Иванов И.И."
    assert out["reviewed_at"]
    # автоматический вердикт не тронут
    assert out["quality_class"] == row["quality_class"]
    assert out["violation_type"] == row["violation_type"]


def test_correction_replaces_the_verdict_only_in_its_own_fields(rb_client: TestClient, processed) -> None:
    sid, rows = processed
    row = next(r for r in rows if "позвоночника" in (r["anatomical_region"] or ""))
    r = rb_client.post(
        f"/api/v1/studies/{sid}/images/{row['image_id']}/review",
        json={
            "action": "correct",
            "violation_type": ["Не выравнена ось позвоночника"],
            "comment": "сколиоз",
        },
    )
    assert r.status_code == 200
    out = r.json()["row"]
    assert out["review_status"] == "corrected"
    assert out["reviewed_quality_class"] == "1"
    assert out["reviewed_violation_type"] == "Не выравнена ось позвоночника"
    assert out["review_comment"] == "сколиоз"
    assert out["quality_class"] == row["quality_class"]  # автоматический вердикт на месте


def test_correction_to_no_violations_sets_quality_zero(rb_client: TestClient, processed) -> None:
    sid, rows = processed
    row = rows[0]
    out = rb_client.post(
        f"/api/v1/studies/{sid}/images/{row['image_id']}/review",
        json={"action": "correct", "violation_type": []},
    ).json()["row"]
    assert out["reviewed_quality_class"] == "0"
    assert out["reviewed_violation_type"] == ""


def test_violation_from_another_region_is_rejected(rb_client: TestClient, processed) -> None:
    """Врач ограничен тем же закрытым списком, что и модель."""
    sid, rows = processed
    spine = next(r for r in rows if "позвоночника" in (r["anatomical_region"] or ""))
    r = rb_client.post(
        f"/api/v1/studies/{sid}/images/{spine['image_id']}/review",
        json={"action": "correct", "violation_type": ["Некорректная область интереса"]},
    )
    assert r.status_code == 422
    assert "нет в списке" in r.json()["detail"]


def test_invented_violation_is_rejected(rb_client: TestClient, processed) -> None:
    sid, rows = processed
    r = rb_client.post(
        f"/api/v1/studies/{sid}/images/{rows[0]['image_id']}/review",
        json={"action": "correct", "violation_type": ["Что-то своё"]},
    )
    assert r.status_code == 422


def test_review_can_be_undone(rb_client: TestClient, processed) -> None:
    sid, rows = processed
    image_id = rows[0]["image_id"]
    rb_client.post(f"/api/v1/studies/{sid}/images/{image_id}/review", json={"action": "confirm"})
    out = rb_client.post(f"/api/v1/studies/{sid}/images/{image_id}/review", json={"action": "reset"}).json()["row"]
    assert out["review_status"] == ""
    assert out["reviewed_quality_class"] is None
    assert out["reviewed_at"] is None


def test_review_reports_allowed_list_and_agreement(rb_client: TestClient, processed) -> None:
    sid, rows = processed
    body = rb_client.post(
        f"/api/v1/studies/{sid}/images/{rows[0]['image_id']}/review", json={"action": "confirm"}
    ).json()
    assert set(body["allowed_violations"]) <= {
        "Некорректная укладка",
        "Не выравнена ось позвоночника",
        "Присутствуют посторонние предметы",
        "Некорректная область интереса",
    }
    assert body["agreement"]["проверено"] == 1
    assert body["agreement"]["подтверждено"] == 1
    assert body["agreement"]["согласие"] == 1.0


def test_review_of_unknown_image_is_404(rb_client: TestClient, processed) -> None:
    sid, _ = processed
    r = rb_client.post(f"/api/v1/studies/{sid}/images/deadbeef/review", json={"action": "confirm"})
    assert r.status_code == 404


# --- выгрузка: автоматический вердикт и решение врача --------------------------


def test_export_defaults_to_the_automatic_verdict(rb_client: TestClient, processed) -> None:
    """В файле, по которому оценивают модель, правок человека быть не должно."""
    sid, rows = processed
    row = next(r for r in rows if "позвоночника" in (r["anatomical_region"] or ""))
    rb_client.post(
        f"/api/v1/studies/{sid}/images/{row['image_id']}/review",
        json={"action": "correct", "violation_type": ["Присутствуют посторонние предметы"]},
    )
    auto = rb_client.get(f"/api/v1/studies/{sid}/download?format=csv").content.decode("utf-8")
    reviewed = rb_client.get(f"/api/v1/studies/{sid}/download?format=csv&source=reviewed").content.decode("utf-8")

    auto_line = next(line for line in auto.splitlines()[1:] if row["image_uid"] in line)
    rev_line = next(line for line in reviewed.splitlines()[1:] if row["image_uid"] in line)
    assert (row["violation_type"] or "") in auto_line
    assert "Присутствуют посторонние предметы" in rev_line
    assert auto_line != rev_line or (row["violation_type"] == "Присутствуют посторонние предметы")


def test_xlsx_has_a_review_sheet(rb_client: TestClient, processed) -> None:
    sid, rows = processed
    rb_client.post(
        f"/api/v1/studies/{sid}/images/{rows[0]['image_id']}/review",
        json={"action": "confirm", "reviewed_by": "Петров"},
    )
    content = rb_client.get(f"/api/v1/studies/{sid}/download?format=xlsx").content
    wb = load_workbook(io.BytesIO(content))
    assert "review" in wb.sheetnames
    ws = wb["review"]
    header = [c.value for c in ws[1]]
    assert "quality_class (сервис)" in header and "quality_class (врач)" in header
    values = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
    assert any(r[header.index("проверил")] == "Петров" for r in values)
    assert any(r[header.index("проверка")] == "подтверждено" for r in values)


def test_sr_records_the_specialist_decision(rb_client: TestClient, processed) -> None:
    sid, rows = processed
    rb_client.post(
        f"/api/v1/studies/{sid}/images/{rows[0]['image_id']}/review",
        json={"action": "confirm", "reviewed_by": "Иванов"},
    )
    ds = pydicom.dcmread(io.BytesIO(rb_client.get(f"/api/v1/studies/{sid}/sr").content))
    assert ds.VerificationFlag == "VERIFIED"
    texts = _texts(ds)
    assert texts["Проверка специалистом"] == "подтверждено специалистом"
    assert texts["Проверил"] == "Иванов"


def test_verified_sr_names_the_verifier(rb_client: TestClient, processed) -> None:
    """VERIFIED без VerifyingObserverSequence — нарушение стандарта (атрибут 1C)."""
    sid, rows = processed
    rb_client.post(
        f"/api/v1/studies/{sid}/images/{rows[0]['image_id']}/review",
        json={"action": "confirm", "reviewed_by": "Иванов"},
    )
    ds = pydicom.dcmread(io.BytesIO(rb_client.get(f"/api/v1/studies/{sid}/sr").content))
    assert ds.VerificationFlag == "VERIFIED"
    assert [str(o.VerifyingObserverName) for o in ds.VerifyingObserverSequence] == ["Иванов"]
    assert len(ds.VerifyingObserverSequence[0].VerificationDateTime) == 14


def test_review_without_a_name_does_not_verify_sr(rb_client: TestClient, processed) -> None:
    """Исправление без имени проверившего — не заверенный документ."""
    sid, rows = processed
    rb_client.post(
        f"/api/v1/studies/{sid}/images/{rows[0]['image_id']}/review",
        json={"action": "correct", "violation_type": []},
    )
    ds = pydicom.dcmread(io.BytesIO(rb_client.get(f"/api/v1/studies/{sid}/sr").content))
    assert ds.VerificationFlag == "UNVERIFIED"
    assert "VerifyingObserverSequence" not in ds
    assert _texts(ds)["Проверка специалистом"]  # само решение в отчёте есть


def test_sr_has_all_type2_attributes(rb_client: TestClient, processed) -> None:
    """Атрибуты типа 2 обязаны присутствовать, хотя бы пустыми (так проверяет dsrdump)."""
    sid, _ = processed
    ds = pydicom.dcmread(io.BytesIO(rb_client.get(f"/api/v1/studies/{sid}/sr").content))
    for tag in (
        "PatientName",
        "PatientID",
        "PatientBirthDate",
        "PatientSex",
        "StudyDate",
        "StudyTime",
        "ReferringPhysicianName",
        "AccessionNumber",
        "ReferencedPerformedProcedureStepSequence",
        "PerformedProcedureCodeSequence",
    ):
        assert tag in ds, tag


def test_sr_explains_non_standard_and_failed_images(rb_client: TestClient, samples: Path) -> None:
    """У нестандартного снимка и у отказа в SR нет пустых TEXT (TextValue — тип 1)."""
    for name in ("tiny.dcm", "not_a_dicom.dcm"):
        f = samples / "edge_cases" / name
        sid = upload(rb_client, (f, f.name)).json()["studies"][0]["id"]
        rb_client.post(f"/api/v1/studies/{sid}/process")
        wait_done(rb_client, sid)
        ds = pydicom.dcmread(io.BytesIO(rb_client.get(f"/api/v1/studies/{sid}/sr").content))
        texts = _texts(ds)
        assert all(v.strip() for v in texts.values()), texts
        assert texts["Анатомическая область"] == "не определена"
        assert texts["Класс качества (0 — годно, 1 — нарушение)"] == "не оценивался"
        if name == "tiny.dcm":
            assert "маленький" in texts["Нестандартные данные"]
            assert texts["Статус обработки"] == "Success"
        else:
            assert texts["Статус обработки"] == "Failure"
