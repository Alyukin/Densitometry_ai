"""Сквозной путь через API на процессоре с правилами ТЗ."""

from __future__ import annotations

import csv
import io
from pathlib import Path

from fastapi.testclient import TestClient

from tests.conftest import upload, wait_done

SPINE_VIOLATIONS = {"Некорректная укладка", "Не выравнена ось позвоночника", "Присутствуют посторонние предметы"}
FEMUR_VIOLATIONS = {"Некорректная укладка", "Некорректная область интереса"}
REGIONS = {"Поясничный отдел позвоночника", "Проксимальный отдел бедра"}


def test_health_reports_rulebased(rb_client: TestClient) -> None:
    body = rb_client.get("/health").json()
    assert body["processor"] == "rulebased"
    assert body["is_mock"] is False


def test_processing_gives_tz_values_and_explanation(rb_client: TestClient, samples: Path) -> None:
    files = sorted((samples / "study_combined").glob("*.dcm"))
    r = upload(rb_client, *[(f, f.name) for f in files])
    assert r.status_code == 201
    sid = r.json()["studies"][0]["id"]

    rb_client.post(f"/api/v1/studies/{sid}/process")
    assert wait_done(rb_client, sid)["status"] == "completed"

    res = rb_client.get(f"/api/v1/studies/{sid}/result").json()
    assert res["is_mock"] is False
    assert len(res["rows"]) == len(files)
    for row in res["rows"]:
        assert row["processing_status"] == "success"
        assert row["anatomical_region"] in REGIONS
        assert row["quality_class"] in ("0", "1")
        allowed = SPINE_VIOLATIONS if "позвоночника" in row["anatomical_region"] else FEMUR_VIOLATIONS
        for v in filter(None, (row["violation_type"] or "").split(";")):
            assert v in allowed, v
        assert 0.0 <= row["confidence"] <= 1.0
        details = row["details"]
        assert details["explanation"]
        assert details["checks"]
        for c in details["checks"]:
            # каждое срабатывание объяснимо: величина, порог и источник критерия
            assert c["measured"] and c["criterion"] and c["source"] in ("ТЗ", "разметка")
        assert details["measurements"]
        assert details["pixel_spacing_mm"] == {"y": 1.05, "x": 0.6}


def test_csv_export_has_quality_prob(rb_client: TestClient, samples: Path) -> None:
    f = next((samples / "study_spine_01").glob("*.dcm"))
    sid = upload(rb_client, (f, f.name)).json()["studies"][0]["id"]
    rb_client.post(f"/api/v1/studies/{sid}/process")
    wait_done(rb_client, sid)

    text = rb_client.get(f"/api/v1/studies/{sid}/download?format=csv").content.decode("utf-8")
    rows = list(csv.DictReader(io.StringIO(text)))
    assert rows
    head = list(rows[0])
    assert head[:8] == [
        "path_to_study",
        "study_uid",
        "image_uid",
        "anatomical_region",
        "quality_class",
        "violation_type",
        "processing_status",
        "time_of_processing",
    ]
    assert head[8] == "quality_prob"
    assert 0.0 <= float(rows[0]["quality_prob"]) <= 1.0
    assert rows[0]["quality_class"] in ("0", "1")


def test_overlay_endpoint_returns_png(rb_client: TestClient, samples: Path) -> None:
    f = next((samples / "study_hip_01").glob("*.dcm"))
    sid = upload(rb_client, (f, f.name)).json()["studies"][0]["id"]
    detail = rb_client.get(f"/api/v1/studies/{sid}").json()
    image_id = detail["images"][0]["id"]

    # до обработки разметки нет
    assert rb_client.get(f"/api/v1/studies/{sid}/images/{image_id}/overlay").status_code == 404

    rb_client.post(f"/api/v1/studies/{sid}/process")
    wait_done(rb_client, sid)

    r = rb_client.get(f"/api/v1/studies/{sid}/images/{image_id}/overlay")
    assert r.status_code == 200
    assert r.headers["content-type"] == "image/png"
    assert r.content[:8] == b"\x89PNG\r\n\x1a\n"


def test_broken_file_is_reported_not_crashed(rb_client: TestClient, samples: Path) -> None:
    f = samples / "edge_cases" / "not_a_dicom.dcm"
    r = upload(rb_client, (f, f.name))
    sid = r.json()["studies"][0]["id"]
    rb_client.post(f"/api/v1/studies/{sid}/process")
    wait_done(rb_client, sid)
    res = rb_client.get(f"/api/v1/studies/{sid}/result").json()
    row = res["rows"][0]
    assert row["processing_status"] == "error"
    assert row["error_message"]


def test_results_are_reproducible(rb_client: TestClient, samples: Path) -> None:
    """ТЗ требует воспроизводимости: одинаковый вход — одинаковый результат."""
    f = next((samples / "study_hip_02_rotated").glob("*.dcm"))
    out = []
    for _ in range(2):
        sid = upload(rb_client, (f, f.name)).json()["studies"][0]["id"]
        rb_client.post(f"/api/v1/studies/{sid}/process")
        wait_done(rb_client, sid)
        row = rb_client.get(f"/api/v1/studies/{sid}/result").json()["rows"][0]
        out.append((row["quality_class"], row["violation_type"], round(row["confidence"], 6)))
        rb_client.delete(f"/api/v1/studies/{sid}")
    assert out[0] == out[1]


def test_export_matches_tz_section_2_5(rb_client: TestClient, samples: Path) -> None:
    """П. 2.5 ТЗ: quality_class — Integer 0/1, processing_status — Success / Failure."""
    good = next((samples / "study_spine_02_tilted").glob("*.dcm"))
    bad = samples / "edge_cases" / "not_a_dicom.dcm"
    sid = upload(rb_client, (good, good.name)).json()["studies"][0]["id"]
    sid_bad = upload(rb_client, (bad, bad.name)).json()["studies"][0]["id"]
    for s in (sid, sid_bad):
        rb_client.post(f"/api/v1/studies/{s}/process")
        wait_done(rb_client, s)

    text = rb_client.get("/api/v1/batch/download?format=csv").content.decode("utf-8")
    rows = list(csv.DictReader(io.StringIO(text)))
    assert rows
    statuses = {r["processing_status"] for r in rows}
    assert statuses <= {"Success", "Failure"}, statuses
    assert "Success" in statuses and "Failure" in statuses
    for r in rows:
        if r["processing_status"] == "Success":
            assert r["quality_class"] in ("0", "1")
            assert int(r["quality_class"]) in (0, 1)

    # в XLSX quality_class должен быть числом, а не текстом
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(rb_client.get("/api/v1/batch/download?format=xlsx").content))
    ws = wb["results"]
    head = [c.value for c in ws[1]]
    col_q = head.index("quality_class") + 1
    col_s = head.index("processing_status") + 1
    for row in range(2, ws.max_row + 1):
        status = ws.cell(row, col_s).value
        assert status in ("Success", "Failure")
        if status == "Success":
            assert isinstance(ws.cell(row, col_q).value, int)
    assert "checks" in wb.sheetnames


def test_checks_sheet_shows_tz_verdict(rb_client: TestClient, samples: Path) -> None:
    f = next((samples / "study_spine_02_tilted").glob("*.dcm"))
    sid = upload(rb_client, (f, f.name)).json()["studies"][0]["id"]
    rb_client.post(f"/api/v1/studies/{sid}/process")
    wait_done(rb_client, sid)

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(rb_client.get(f"/api/v1/studies/{sid}/download?format=xlsx").content))
    ws = wb["checks"]
    head = [c.value for c in ws[1]]
    assert "порог ТЗ" in head and "по букве ТЗ" in head and "участвует в вердикте" in head
    col_check = head.index("check") + 1
    names = {ws.cell(r, col_check).value for r in range(2, ws.max_row + 1)}
    assert "spine_axis_tz" in names  # методика ТЗ считается и попадает в отчёт


def test_tiny_image_is_non_standard(rb_client: TestClient, samples: Path) -> None:
    """Кадр 24×24 — не снимок денситометра: Success без области и класса, а не вердикт."""
    f = samples / "edge_cases" / "tiny.dcm"
    sid = upload(rb_client, (f, f.name)).json()["studies"][0]["id"]
    rb_client.post(f"/api/v1/studies/{sid}/process")
    wait_done(rb_client, sid)
    row = rb_client.get(f"/api/v1/studies/{sid}/result").json()["rows"][0]
    assert row["processing_status"] == "success"
    assert row["anatomical_region"] is None and row["quality_class"] is None
    assert "маленький" in row["details"]["non_standard"]

    text = rb_client.get(f"/api/v1/studies/{sid}/download?format=csv").content.decode("utf-8")
    exported = next(csv.DictReader(io.StringIO(text)))
    assert exported["processing_status"] == "Success"
    assert exported["anatomical_region"] == exported["quality_class"] == exported["quality_prob"] == ""
